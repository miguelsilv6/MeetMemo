"""
DOCX generation utilities for MeetMemo exports.

This module provides functions for generating Word (.docx) transcript
documents in a plain dialogue format, using python-docx.
"""
from datetime import datetime
from io import BytesIO

from config import Settings, get_settings
from docx import Document
from docx.shared import RGBColor

from utils.formatters import format_speaker_name


def generate_transcript_docx(
    meeting_title: str,
    transcript_data: list,
    generated_on: str = None,
    settings: Settings = None
) -> BytesIO:
    """
    Generate a Word (.docx) transcript-only document (no AI summary).

    Each segment is rendered as a single dialogue line in chronological
    order, formatted as "Speaker: text" (no timestamps), e.g.:

        Speaker 1: Ola
        Speaker 2: Boa tarde

    Args:
        meeting_title: Meeting title/filename
        transcript_data: List of transcript segments
        generated_on: Optional formatted timestamp string
        settings: Optional Settings instance for timezone

    Returns:
        BytesIO buffer containing the DOCX file

    Example:
        >>> transcript = [
        ...     {'speaker': 'SPEAKER_00', 'text': 'Hello', 'start': '0.00', 'end': '1.00'}
        ... ]
        >>> docx_buffer = generate_transcript_docx('Team Meeting', transcript)
    """
    if settings is None:
        settings = get_settings()

    if not generated_on:
        generated_on = datetime.now(settings.timezone).strftime('%B %d, %Y at %I:%M %p')

    document = Document()

    document.add_heading(meeting_title or 'MeetMemo Meeting Transcript', level=1)

    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(f"Generated on {generated_on}")
    subtitle_run.italic = True
    subtitle_run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    for entry in transcript_data:
        text = (entry.get('text') or '').strip()
        if not text:
            continue

        speaker = format_speaker_name(entry.get('speaker', 'Unknown Speaker'))
        paragraph = document.add_paragraph()
        speaker_run = paragraph.add_run(f"{speaker}: ")
        speaker_run.bold = True
        paragraph.add_run(text)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
